"""DOCX extraction tests with generated, synthetic Word documents."""

from io import BytesIO

from app.core.exceptions import ExtractionErrorCode
from app.extractors.docx import DOCXExtractor
from app.models.case import Case
from app.models.document import Document, DocumentStatus
from app.services.document_service import DocumentService
from docx import Document as WordDocument
from tests.fakes import InMemoryRepository, InMemoryStorage


def _document() -> Document:
    return Document(
        document_id="docx-test-document",
        case_id="docx-test-case",
        filename="synthetic.docx",
        mime_type=(
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        ),
        size=1,
        storage_path="synthetic/object",
    )


def _docx_bytes(with_text: bool = True) -> bytes:
    source = WordDocument()
    if with_text:
        source.add_paragraph("First synthetic paragraph")
        table = source.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Middle table cell A"
        table.cell(0, 1).text = "Middle table cell B"
        source.add_paragraph("Last synthetic paragraph")
    stream = BytesIO()
    source.save(stream)
    return stream.getvalue()


def test_docx_extracts_paragraphs_and_tables_in_document_order() -> None:
    result = DOCXExtractor().extract(_document(), _docx_bytes())

    assert result.status is DocumentStatus.SUCCESS
    assert result.extraction_method == "docx"
    assert len(result.pages) == 1
    page = result.pages[0]
    assert page.extraction_method == "docx"
    assert page.page_number == 1
    assert page.word_boxes == []
    assert (
        page.text.index("First synthetic paragraph")
        < page.text.index("Middle table cell A")
        < page.text.index("Last synthetic paragraph")
    )


def test_empty_docx_returns_no_text_error() -> None:
    result = DOCXExtractor().extract(_document(), _docx_bytes(with_text=False))

    assert result.status is DocumentStatus.FAILED
    assert result.error is not None
    assert result.error.code is ExtractionErrorCode.NO_TEXT_DETECTED


def test_corrupt_and_zero_byte_docx_return_structured_errors() -> None:
    extractor = DOCXExtractor()
    corrupt = extractor.extract(_document(), b"not a zip document")
    empty = extractor.extract(_document(), b"")

    assert corrupt.status is DocumentStatus.FAILED
    assert corrupt.error is not None
    assert corrupt.error.code is ExtractionErrorCode.DOCX_PARSE_FAILED
    assert empty.status is DocumentStatus.FAILED
    assert empty.error is not None
    assert empty.error.code is ExtractionErrorCode.NO_TEXT_DETECTED


def test_document_service_processes_docx_to_page_one() -> None:
    repository = InMemoryRepository()
    storage = InMemoryStorage()
    case = Case(case_id="docx-service-case")
    repository.save_case(case)
    service = DocumentService(repository, storage)
    document = service.accept_upload(
        case.case_id,
        "synthetic.docx",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        _docx_bytes(),
    )

    completed = service.process_document(case.case_id, document.document_id)

    assert completed is not None
    assert completed.status is DocumentStatus.SUCCESS
    assert completed.extraction_method == "docx"
    page = repository.get_page(case.case_id, document.document_id, 1)
    assert page is not None
    assert page.extraction_method == "docx"
    assert page.word_boxes == []
